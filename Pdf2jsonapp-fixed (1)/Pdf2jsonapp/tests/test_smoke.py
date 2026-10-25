import io
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))

from docx import Document

from extractors.docx_extractor import extract_docx
from utils.json_builder import build_structured_json


def _page(n, *blocks):
    return {"page_num": n, "blocks": [{"text": t, "font_size": s} for t, s in blocks], "ocr_used": False}


def test_builder_chapters_and_same_page_split():
    pages = [
        _page(1, ("Chapter 1", 24), ("Intro", 24), ("body one " * 5, 12), ("body two " * 5, 12)),
        _page(2, ("more body " * 5, 12), ("Chapter 2", 24), ("text under chapter two " * 3, 12)),
    ]
    out = build_structured_json(pages, title="T")
    assert out["chapter_count"] == 2
    assert out["chapters"][0]["heading"] == "Chapter 1 Intro"
    assert len(out["chapters"][0]["pages"]) == 2
    assert "chapter two" in out["chapters"][1]["pages"][0]["text"]
    assert "more body" not in out["chapters"][1]["pages"][0]["text"]


def test_builder_drops_running_headers_and_page_numbers():
    words = ["apples", "rivers", "granite", "violins", "planets", "tunnels", "markets", "glaciers"]
    pages = [
        _page(i, ("MY BOOK TITLE", 20), (f"{i}", 10), (f"{w} is the topic discussed here in detail " * 2, 12))
        for i, w in enumerate(words, start=1)
    ]
    out = build_structured_json(pages)
    assert out["chapter_count"] == 1
    assert all("MY BOOK TITLE" not in p["text"] for p in out["chapters"][0]["pages"])


def test_docx_headings_and_page_breaks():
    d = Document()
    d.add_heading("Chapter One", 1)
    d.add_paragraph("Body of chapter one.")
    d.add_page_break()
    d.add_paragraph("Still chapter one, page two.")
    d.add_heading("Chapter Two", 1)
    d.add_paragraph("Body of chapter two.")
    buf = io.BytesIO()
    d.save(buf)
    pages = extract_docx(buf.getvalue())
    assert len(pages) == 2
    out = build_structured_json(pages, title="doc")
    assert [c["heading"] for c in out["chapters"]] == ["Chapter One", "Chapter Two"]
